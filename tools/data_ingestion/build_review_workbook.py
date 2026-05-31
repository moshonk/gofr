#!/usr/bin/env python3
"""
Compress the facility data-quality task into a single, navigable Excel workbook
for a reviewer.

The workbook contains:
  * Overview      — hyperlinked contents + headline metrics.
  * Read Me       — gap descriptions, severity legend, caveats, methodology.
  * Summary       — facilities-per-province and facilities-per-type tables.
  * Combined Data — the full merged dataset (datasets/combined.csv).
  * One tab per issue line-list / summary table from datasets/data_quality_issues/.

Sources (regenerate these first if stale):
  python3 tools/data_ingestion/merge_facility_master_list.py
  python3 tools/data_ingestion/generate_hfids.py
  python3 tools/data_ingestion/analyze_facility_data_quality.py
  python3 tools/data_ingestion/export_data_quality_issues.py

Usage:
  python3 tools/data_ingestion/build_review_workbook.py
"""
import csv
import os
from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
CSV_PATH = os.path.join(REPO_ROOT, "datasets", "combined.csv")
ISSUE_DIR = os.path.join(REPO_ROOT, "datasets", "data_quality_issues")
INDEX_CSV = os.path.join(ISSUE_DIR, "_index.csv")
OUT_PATH = os.path.join(REPO_ROOT, "datasets", "CamDHEA-facility-data-quality-review.xlsx")

# Short, Excel-safe (<=31 char) tab names keyed by issue_id, in review order.
SHEET_NAMES = {
    "completeness_missing_required": "01 Missing required",
    "validity_code_format": "02 Code format",
    "validity_hierarchy_nesting": "03 Hierarchy nesting",
    "validity_gps": "04 GPS",
    "validity_value_set": "05 Value-set typos",
    "validity_placeholder_email": "06 Placeholder email",
    "validity_unparseable_start_date": "07 Bad start date",
    "duplicate_hf_code": "08 Duplicate HF_CODE",
    "duplicate_name_within_commune": "09 Dup name in commune",
    "coverage_districts_no_facilities": "10 Districts no facility",
    "coverage_communes_no_facilities": "11 Communes no facility",
    "coverage_orphan_district_codes": "12 Orphan district codes",
    "coverage_orphan_commune_codes": "13 Orphan commune codes",
    "integrity_province_code_name_conflict": "14 Province name conflict",
    "integrity_district_code_name_conflict": "15 District name conflict",
    "integrity_commune_code_name_conflict": "16 Commune name conflict",
    "dupnames_reused_commune_names": "17 Reused commune names",
    "dupnames_facility_equals_commune": "18 Facility=commune name",
    "dupnames_reused_facility_names": "19 Reused facility names",
}
SEV_TAB_COLOR = {"Blocker": "C00000", "Major": "ED7D31", "Minor": "FFC000"}
SEV_FILL = {
    "Blocker": PatternFill("solid", fgColor="F4CCCC"),
    "Major": PatternFill("solid", fgColor="FCE4D6"),
    "Minor": PatternFill("solid", fgColor="FFF2CC"),
}

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        return list(reader)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def autosize(ws, rows, header_row, ncols, cap=55):
    for c in range(1, ncols + 1):
        width = 10
        for r in rows:
            if c - 1 < len(r):
                width = max(width, min(cap, len(str(r[c - 1])) + 2))
        ws.column_dimensions[get_column_letter(c)].width = width


def back_link(ws):
    cell = ws.cell(row=1, column=1, value="« Back to Overview")
    cell.hyperlink = "#'Overview'!A1"
    cell.font = LINK_FONT


def write_data_sheet(ws, title, description, table, tab_color=None):
    """Generic sheet: back-link, title, description, header row, data, filter."""
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    back_link(ws)
    ws.cell(row=2, column=1, value=title).font = TITLE_FONT
    if description:
        d = ws.cell(row=3, column=1, value=description)
        d.alignment = WRAP
    header_row = 5
    if not table:
        ws.cell(row=header_row, column=1, value="(no records — this issue did not occur)")
        return
    ncols = max(len(r) for r in table)
    for j, val in enumerate(table[0], start=1):
        ws.cell(row=header_row, column=j, value=val)
    style_header(ws, header_row, ncols)
    for i, row in enumerate(table[1:], start=header_row + 1):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{header_row + len(table) - 1}"
    autosize(ws, table, header_row, ncols)


def main():
    if not os.path.isfile(INDEX_CSV):
        raise SystemExit(
            "Missing datasets/data_quality_issues/_index.csv — run "
            "export_data_quality_issues.py first.")

    combined = read_csv(CSV_PATH)
    index = read_csv(INDEX_CSV)
    index_header, index_rows = index[0], index[1:]

    wb = Workbook()

    # ---------------- Overview ----------------
    ov = wb.active
    ov.title = "Overview"
    ov.sheet_properties.tabColor = "305496"
    ov.cell(row=1, column=1,
            value="CamDHEA Health Facility Master List — Data Quality Review").font = Font(bold=True, size=16)
    ov.cell(row=2, column=1,
            value=f"Generated {datetime.now():%Y-%m-%d %H:%M}  ·  "
                  f"{len(combined) - 1} facility records  ·  source: datasets/combined.csv")

    sev_counts = Counter(r[2] for r in index_rows if int(r[4]) > 0)
    ov.cell(row=4, column=1, value="Issues flagged:").font = Font(bold=True)
    ov.cell(row=4, column=2,
            value=f"{sev_counts.get('Blocker', 0)} Blocker · "
                  f"{sev_counts.get('Major', 0)} Major · "
                  f"{sev_counts.get('Minor', 0)} Minor")

    # Contents table
    hr = 6
    headers = ["#", "Go", "Issue dataset", "Category", "Severity", "Records", "Description"]
    for j, h in enumerate(headers, start=1):
        ov.cell(row=hr, column=j, value=h)
    style_header(ov, hr, len(headers))

    r = hr + 1
    # Fixed informational sheets first
    for n, (sheet, desc) in enumerate([
        ("Read Me", "Gap descriptions, severity legend and caveats"),
        ("Summary", "Facilities per province and per facility type"),
        ("Combined Data", "Full merged dataset (all facilities, canonical schema)"),
    ], start=1):
        ov.cell(row=r, column=1, value="")
        link = ov.cell(row=r, column=2, value="open")
        link.hyperlink = f"#'{sheet}'!A1"
        link.font = LINK_FONT
        ov.cell(row=r, column=3, value=sheet)
        ov.cell(row=r, column=7, value=desc)
        r += 1

    for i, row in enumerate(index_rows, start=1):
        issue_id, category, severity, description, count, _file = row
        sheet = SHEET_NAMES.get(issue_id, issue_id[:31])
        ov.cell(row=r, column=1, value=i)
        if int(count) > 0:
            link = ov.cell(row=r, column=2, value="open")
            link.hyperlink = f"#'{sheet}'!A1"
            link.font = LINK_FONT
        else:
            ov.cell(row=r, column=2, value="—")
        ov.cell(row=r, column=3, value=sheet)
        ov.cell(row=r, column=4, value=category)
        sev_cell = ov.cell(row=r, column=5, value=severity)
        if severity in SEV_FILL:
            sev_cell.fill = SEV_FILL[severity]
        ov.cell(row=r, column=6, value=int(count))
        ov.cell(row=r, column=7, value=description).alignment = WRAP
        r += 1

    ov.freeze_panes = "A7"
    widths = [4, 6, 26, 16, 10, 9, 60]
    for j, wdt in enumerate(widths, start=1):
        ov.column_dimensions[get_column_letter(j)].width = wdt

    # ---------------- Read Me ----------------
    rm = wb.create_sheet("Read Me")
    rm.sheet_properties.tabColor = "305496"
    back_link(rm)
    lines = [
        ("This workbook", "Single-file package of the CamDHEA facility data-quality task: the merged national Health Facility Master List, a summary, and a separate line-list for every gap found in the data-quality report."),
        ("How to navigate", "Start on the Overview tab and click 'open' to jump to any sheet. Each sheet has a '« Back to Overview' link in cell A1. Issue sheets have filter dropdowns on the header row."),
        ("", ""),
        ("Severity — Blocker", "Breaks or corrupts FHIR ingestion; must be fixed before loading (duplicate HF_CODE, hierarchy-nesting violations)."),
        ("Severity — Major", "Registry materially degraded: missing required fields, missing/invalid GPS, coverage gaps, orphan admin codes, code↔name conflicts."),
        ("Severity — Minor", "Cosmetic or easily normalised: value-set typos, placeholder emails, reused names (many legitimately valid)."),
        ("", ""),
        ("Diagnostic columns", "Facility line-lists keep the full canonical schema plus a diagnostic column (missing_required_fields, code_issues, nesting_issues, gps_issue, value_issues, date_issue) explaining the defect."),
        ("Caveat — completeness", "'Missing required' is dominated by HF_LEVEL (blank on most non-CPA facilities) and the placeholder-email count by literal 'None'/'NA' values — high counts but mostly Minor."),
        ("Caveat — references", "Coverage and reused-name tables are expectation references, not always defects: Cambodian commune names repeat legitimately and not every commune hosts a facility. Boundary references are valid as of 2018, so genuine post-2018 admin units appear as 'orphan' codes."),
        ("Caveat — villages", "No national village registry is bundled, so villages-with-no-facilities cannot be enumerated against a full denominator."),
        ("Full report", "docs/data_quality/facility-data-quality-report.md (narrative) and docs/data_quality/facility-data-quality-plan.md (methodology)."),
    ]
    rr = 3
    for label, text in lines:
        if label:
            rm.cell(row=rr, column=1, value=label).font = Font(bold=True)
        rm.cell(row=rr, column=2, value=text).alignment = WRAP
        rr += 1
    rm.column_dimensions["A"].width = 22
    rm.column_dimensions["B"].width = 110

    # ---------------- Summary ----------------
    sm = wb.create_sheet("Summary")
    sm.sheet_properties.tabColor = "305496"
    data_rows = combined[1:]
    hidx = {h: i for i, h in enumerate(combined[0])}
    prov = Counter((r[hidx["PRO_NAME_EN"]] or "?").strip() for r in data_rows)
    typ = Counter((r[hidx["HF_TYPE_EN"]] or "(blank)").strip() for r in data_rows)
    back_link(sm)
    sm.cell(row=2, column=1, value="Dataset summary").font = TITLE_FONT
    sm.cell(row=4, column=1, value="Facilities per province").font = Font(bold=True)
    sm.cell(row=5, column=1, value="Province"); sm.cell(row=5, column=2, value="Facilities")
    style_header(sm, 5, 2)
    r = 6
    for p, c in prov.most_common():
        sm.cell(row=r, column=1, value=p); sm.cell(row=r, column=2, value=c); r += 1
    start_t = r + 2
    sm.cell(row=start_t, column=1, value="Facilities per type").font = Font(bold=True)
    sm.cell(row=start_t + 1, column=1, value="Type"); sm.cell(row=start_t + 1, column=2, value="Facilities")
    style_header(sm, start_t + 1, 2)
    r = start_t + 2
    for t, c in typ.most_common():
        sm.cell(row=r, column=1, value=t); sm.cell(row=r, column=2, value=c); r += 1
    sm.column_dimensions["A"].width = 28
    sm.column_dimensions["B"].width = 14

    # ---------------- Combined Data ----------------
    cd = wb.create_sheet("Combined Data")
    write_data_sheet(cd, "Combined Dataset — all facilities",
                     f"{len(data_rows)} facility records · canonical 32-column schema · "
                     "merged from 25 provincial master-list workbooks.",
                     combined, tab_color="2E7D32")

    # ---------------- Issue sheets ----------------
    for row in index_rows:
        issue_id, category, severity, description, count, fname = row
        if int(count) == 0 or not fname:
            continue
        table = read_csv(os.path.join(ISSUE_DIR, fname))
        sheet = SHEET_NAMES.get(issue_id, issue_id[:31])
        ws = wb.create_sheet(sheet)
        write_data_sheet(
            ws,
            f"{sheet}  ({severity})",
            f"{description}.  Records: {count}.",
            table,
            tab_color=SEV_TAB_COLOR.get(severity))

    wb.save(OUT_PATH)
    print(f"Wrote workbook -> {os.path.relpath(OUT_PATH, REPO_ROOT)} "
          f"({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
